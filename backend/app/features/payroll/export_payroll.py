from io import BytesIO

from fastapi.responses import StreamingResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)
from openpyxl.utils import get_column_letter

from app.db.mongo import get_db_conn
from app.db.collections import collections
from app.logging_config import logger


async def export_salary_sheet(
    month: int,
    year: int
):
    try:

        db = await get_db_conn()

        payrolls = await db[
            collections.PAYROLL
        ].find(
            {
                "month": month,
                "year": year
            }
        ).to_list(length=None)

        if not payrolls:
            raise ValueError(
                "No payroll records found"
            )

        payroll_summary = await db[
            collections.PAYROLL_SUMMARY
        ].find_one(
            {
                "month": month,
                "year": year
            }
        )

        if not payroll_summary:
            raise ValueError(
                "Payroll summary not found"
            )

        # =========================
        # Create Workbook
        # =========================

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "Payroll"

        # =========================
        # Headers
        # =========================

        headers = [
            "Employee No",
            "Employee Name",
            "Department",
            "Designation",
            "Employment Type",
            "Monthly Salary",
            "Per Day Salary",
            "Working Days",
            "Worked Days",
            "Worked Days Override",
            "Absent Days",
            "Monthly Leave Used",
            "Annual Leave Used",
            "Deductible Days",
            "Leave Deduction",
            "Late Count",
            "Late Deduction",
            "Half Days",
            "Half Day Deduction",
            "Half Days Override",
            "Bonus",
            "Loan Deduction",
            "Total Deductions",
            "Final Salary",
            "Salary Received"
        ]

        worksheet.append(headers)

        # =========================
        # Styles
        # =========================

        header_fill = PatternFill(
            start_color="1F4E78",
            end_color="1F4E78",
            fill_type="solid"
        )

        header_font = Font(
            color="FFFFFF",
            bold=True
        )

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        center_alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        # =========================
        # Header Styling
        # =========================

        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin_border
            cell.alignment = center_alignment

        # Freeze Header

        worksheet.freeze_panes = "A2"

        # =========================
        # Employee Payroll Records
        # =========================

        for payroll in payrolls:

            total_deductions = (
                payroll.get(
                    "leave_deduction",
                    0
                )
                + payroll.get(
                    "late_deduction",
                    0
                )
                + payroll.get(
                    "half_day_deduction",
                    0
                )
                + payroll.get(
                    "loan_deduction",
                    0
                )
            )

            worksheet.append(
                [
                    payroll.get("employee_no"),
                    payroll.get("employee_name"),
                    payroll.get("department"),
                    payroll.get("designation"),
                    payroll.get("employment_type"),
                    payroll.get("monthly_salary"),
                    payroll.get("per_day_salary"),
                    payroll.get("working_days"),
                    payroll.get("worked_days"),
                    payroll.get("worked_days_override"),
                    payroll.get("absent_days"),
                    payroll.get("monthly_leave_used"),
                    payroll.get("annual_leave_used"),
                    payroll.get("deductible_days"),
                    payroll.get("leave_deduction"),
                    payroll.get("late_count"),
                    payroll.get("late_deduction"),
                    payroll.get("half_days"),
                    payroll.get("half_day_deduction"),
                    payroll.get("half_days_override"),
                    payroll.get("bonus"),
                    payroll.get("loan_deduction"),
                    total_deductions,
                    payroll.get("final_salary"),
                    (
                        "Yes"
                        if payroll.get(
                            "salary_received"
                        )
                        else "No"
                    )
                ]
            )

            # =========================
            # Style Current Row
            # =========================

            current_row = worksheet.max_row

            for cell in worksheet[current_row]:

                cell.border = thin_border
                cell.alignment = center_alignment

        # =========================
        # Currency Formatting
        # =========================

        currency_columns = [
            6,   # Monthly Salary
            7,   # Per Day Salary
            15,  # Leave Deduction
            17,  # Late Deduction
            19,  # Half Day Deduction
            21,  # Bonus
            22,  # Loan Deduction
            23,  # Total Deductions
            24   # Final Salary
        ]

        for col in currency_columns:

            for column_cells in worksheet.iter_cols(
                min_col=col,
                max_col=col,
                min_row=2,
                max_row=worksheet.max_row
            ):

                for cell in column_cells:

                    cell.number_format = (
                        '#,##0.00'
                    )

        # =========================
        # Payroll Summary
        # =========================

        worksheet.append([])

        worksheet.append(
            ["Payroll Summary"]
        )

        worksheet.append(
            [
                "Month",
                payroll_summary.get("month")
            ]
        )

        worksheet.append(
            [
                "Year",
                payroll_summary.get("year")
            ]
        )

        worksheet.append(
            [
                "Total Employees Processed",
                payroll_summary.get(
                    "total_employees_processed"
                )
            ]
        )

        worksheet.append(
            [
                "Total Gross Salary",
                payroll_summary.get(
                    "total_gross_salary"
                )
            ]
        )

        worksheet.append(
            [
                "Total Deductions",
                payroll_summary.get(
                    "total_deductions"
                )
            ]
        )

        worksheet.append(
            [
                "Total Net Payroll",
                payroll_summary.get(
                    "total_net_payroll"
                )
            ]
        )

        worksheet.append(
            [
                (
                    "Total Amount Required "
                    "For Salary Disbursement"
                ),
                payroll_summary.get(
                    "total_amount_required_for_salary_disbursement"
                )
            ]
        )

        # =========================
        # Payroll Summary Styling
        # =========================

        summary_start_row = (
            worksheet.max_row - 6
        )

        for row in range(
            summary_start_row,
            worksheet.max_row + 1
        ):

            for cell in worksheet[row]:

                cell.border = thin_border
                cell.alignment = center_alignment

            worksheet.cell(
                row=row,
                column=1
            ).font = Font(
                bold=True
            )

            worksheet.cell(
                row=row,
                column=1
            ).fill = PatternFill(
                start_color="D9EAF7",
                end_color="D9EAF7",
                fill_type="solid"
            )

        # =========================
        # Auto Adjust Column Width
        # =========================

        for column_cells in worksheet.columns:

            length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )

            worksheet.column_dimensions[
                get_column_letter(
                    column_cells[0].column
                )
            ].width = length + 5

        # =========================
        # Save Workbook
        # =========================

        excel_file = BytesIO()

        workbook.save(excel_file)

        excel_file.seek(0)

        return StreamingResponse(
            excel_file,
            media_type=(
                "application/"
                "vnd.openxmlformats-"
                "officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                (
                    f'attachment; '
                    f'filename="Payroll_{month}_{year}.xlsx"'
                )
            }
        )

    except Exception as e:

        logger.exception(
            f"Error exporting salary sheet: "
            f"{str(e)}"
        )

        raise